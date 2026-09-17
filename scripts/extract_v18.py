# -*- coding: utf-8 -*-
"""Extract full text + structure of the v1.8 SLR manuscript for review."""
import re, sys
from docx import Document

P = r'h:\phd\6.pjt\LR1\papers\05_article\SLR for CDT (v1.8 260906).docx'
doc = Document(P)

paras = [(p.style.name, p.text) for p in doc.paragraphs]
n_words = sum(len(re.findall(r'\w+', t)) for _, t in paras)
print('paragraphs:', len(paras), '| words (approx):', n_words)
print('tables:', len(doc.tables))

# headings map
print('\n=== HEADINGS ===')
for s, t in paras:
    if s.startswith('Heading') and t.strip():
        print('  [%s] %s' % (s, t.strip()[:120]))

# abstract + keywords region: first ~40 non-empty paragraphs
print('\n=== FRONT MATTER (first 40 non-empty paras) ===')
c = 0
for s, t in paras:
    if t.strip():
        print('<%s> %s' % (s, t.strip()[:300]))
        c += 1
        if c >= 40:
            break

# write full text for analysis
out = r'h:\phd\6.pjt\LR1\papers\scripts\_v18_fulltext.txt'
with open(out, 'w', encoding='utf-8') as f:
    for s, t in paras:
        if t.strip():
            f.write('<%s> %s\n' % (s, t.rstrip()))
print('\nfull text written to', out)

# quick content probes
full = '\n'.join(t for _, t in paras)
probes = {
    'PRISMA mention': len(re.findall(r'PRISMA', full, re.I)),
    'PRISMA 2020': len(re.findall(r'PRISMA 2020', full)),
    'n= counts': re.findall(r'\bn\s*=\s*[\d,]+', full)[:15],
    'years mentioned': re.findall(r'\b(19|20)\d\d\b', full)[:0],
    'refs section': bool(re.search(r'(References|REFERENCES)', full)),
    'fusion depth': len(re.findall(r'fusion depth', full, re.I)),
    'Tier 4': len(re.findall(r'Tier 4', full)),
    'self-improv': len(re.findall(r'self-improv', full, re.I)),
    'Automation in Construction': len(re.findall(r'Automation in Construction', full)),
    'Figure mentions': len(re.findall(r'\bFig(?:ure)?\.\s*\d+', full)),
    'Table mentions': len(re.findall(r'\bTable\s*\d+', full)),
    'Limitations': len(re.findall(r'imitation', full, re.I)),
    'future work': len(re.findall(r'future (work|research)', full, re.I)),
    'TODO/placeholder': re.findall(r'(?i)(TODO|XXX|\[citation needed\]|\bTBD\b|placeholder|FIXME)', full),
    'errors/proofread marks': re.findall(r'(?i)\b(error|delete this|rewrite)\b', full)[:5],
}
print('\n=== PROBES ===')
for k, v in probes.items():
    print(' ', k, ':', v)
